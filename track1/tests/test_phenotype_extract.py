from __future__ import annotations

import io
from contextlib import redirect_stdout
from pathlib import Path

from docx import Document

from hpo_fixtures import write_mini_ontology
from mva_track1.hpo_ontology import parse_obo
from mva_track1.phenotype_extract import (
    classify_status,
    extract_phenotype,
    format_safe_phenotype_report,
    format_structure_report,
    inspect_docx_structure,
    map_block,
    profile_terms,
    write_private_candidates,
    write_safe_phenotype,
)

SECRET_SENTENCE = "SECRETCLINICALPHRASE_DO_NOT_EMIT"
SECRET_NAME = "SECRETPERSONNAMEZZZ"


def _write_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph(f"HP:0001250 {SECRET_NAME} {SECRET_SENTENCE}")
    doc.add_paragraph("Seizure")
    doc.add_paragraph("No intellectual disability")
    doc.add_paragraph("Possible macrocephaly")
    doc.add_paragraph("Neurological abnormality")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Finding"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Seizures"
    table.cell(1, 1).text = "present"
    doc.save(path)


def test_docx_structure_and_explicit_hpo(tmp_path: Path) -> None:
    ontology = parse_obo(write_mini_ontology(tmp_path))
    docx = tmp_path / "note.docx"
    _write_docx(docx)
    structure = inspect_docx_structure(docx, ontology)
    assert structure.paragraph_count >= 5
    assert structure.table_count == 1
    assert structure.table_shapes[0]["rows"] == 2
    assert structure.table_shapes[0]["cols"] == 2
    assert structure.distinct_explicit_hpo == 1
    assert structure.explicit_hpo_ids == ["HP:0001250"]
    assert structure.appears_structured is True
    text = format_structure_report(structure)
    assert SECRET_SENTENCE not in text
    assert SECRET_NAME not in text


def test_label_synonym_negation_uncertain(tmp_path: Path) -> None:
    ontology = parse_obo(write_mini_ontology(tmp_path))
    present = map_block("Seizure", ontology)
    assert any(
        item.hpo_id == "HP:0001250"
        and item.status == "present"
        and item.mapping_method == "exact_label"
        for item in present
    )
    synonym = map_block("Seizures", ontology)
    assert any(item.hpo_id == "HP:0001250" and item.mapping_method == "synonym" for item in synonym)
    related = map_block("Large head", ontology)
    assert any(item.hpo_id == "HP:0000256" and item.confidence == "medium" for item in related)
    absent = map_block("No intellectual disability", ontology)
    assert classify_status("No intellectual disability") == "absent"
    assert any(item.hpo_id == "HP:0001249" and item.status == "absent" for item in absent)
    mixed = map_block("No seizures and intellectual disability", ontology)
    assert any(item.hpo_id == "HP:0001250" and item.status == "absent" for item in mixed)
    assert any(item.hpo_id == "HP:0001249" and item.status == "present" for item in mixed)
    assert not any(
        item.hpo_id == "HP:0000707" and item.status == "absent" for item in mixed
    )
    family = map_block("Family history of seizures in the mother HP:0001250", ontology)
    assert any(
        item.hpo_id == "HP:0001250" and item.scope == "family_history" for item in family
    )
    uncertain = map_block("Possible macrocephaly", ontology)
    assert any(item.hpo_id == "HP:0000256" and item.status == "uncertain" for item in uncertain)


def test_stdout_omits_source_text(tmp_path: Path) -> None:
    ontology = parse_obo(write_mini_ontology(tmp_path))
    docx = tmp_path / "note.docx"
    _write_docx(docx)
    structure, candidates = extract_phenotype(docx, ontology)
    buf = io.StringIO()
    with redirect_stdout(buf):
        print(format_structure_report(structure), end="")
        print(format_safe_phenotype_report(candidates), end="")
    text = buf.getvalue()
    assert SECRET_SENTENCE not in text
    assert SECRET_NAME not in text
    assert "HP:0001250" in text
    private = tmp_path / "private.tsv"
    safe_tsv = tmp_path / "safe.tsv"
    safe_json = tmp_path / "safe.json"
    write_private_candidates(private, candidates)
    write_safe_phenotype(safe_tsv, safe_json, candidates)
    safe_text = safe_tsv.read_text(encoding="utf-8")
    assert SECRET_SENTENCE not in safe_text
    assert "source_phrase" not in safe_text
    assert SECRET_SENTENCE in private.read_text(encoding="utf-8")
    profile = profile_terms(candidates)
    assert all(item.confidence == "high" and item.status == "present" for item in profile)
    assert all(item.mapping_method != "lexical_candidate" for item in profile)
    assert all(item.scope == "proband" for item in profile)


def test_family_history_excluded_from_ranking_profile(tmp_path: Path) -> None:
    ontology = parse_obo(write_mini_ontology(tmp_path))
    family = map_block("Maternal history of miscarriage; Seizures HP:0001250", ontology)
    assert any(item.scope == "family_history" for item in family)
    profile = profile_terms(family)
    assert all(item.scope == "proband" for item in profile)


def test_no_broad_parent_negative_inference(tmp_path: Path) -> None:
    ontology = parse_obo(write_mini_ontology(tmp_path))
    # Specific negation must not absent the nervous-system parent.
    mapped = map_block("No seizures", ontology)
    assert any(item.hpo_id == "HP:0001250" and item.status == "absent" for item in mapped)
    assert not any(item.hpo_id == "HP:0000707" and item.status == "absent" for item in mapped)
