"""Local clinical DOCX phenotype extraction and HPO mapping.

Raw clinical wording is never printed. Source phrases may be written only to a
private working-directory file that stays outside Git.
"""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from mva_track1.hpo_ontology import (
    GENERIC_HPO_IDS,
    HPO_ID_RE,
    HpoOntology,
    lookup_normalized,
    normalize_label,
)

STATUS_PRESENT = "present"
STATUS_ABSENT = "absent"
STATUS_UNCERTAIN = "uncertain"

SCOPE_PROBAND = "proband"
SCOPE_FAMILY = "family_history"
SCOPE_UNCERTAIN = "uncertain_context"

METHOD_EXPLICIT = "explicit_hpo"
METHOD_EXACT_LABEL = "exact_label"
METHOD_SYNONYM = "synonym"
METHOD_LEXICAL = "lexical_candidate"

CONF_HIGH = "high"
CONF_MEDIUM = "medium"
CONF_LOW = "low"

NEGATION_RE = re.compile(
    r"\b(?:no|not|without|denies|denied|denying|negative(?:\s+for)?|"
    r"absent|lack(?:s|ing)?(?:\s+of)?|never\s+had|no\s+evidence(?:\s+of)?|"
    r"ruled\s+out|rule[sd]?\s+out)\b",
    re.IGNORECASE,
)
UNCERTAIN_RE = re.compile(
    r"\b(?:possible|possibly|probable|probably|suspected|suspect|"
    r"questionable|uncertain|equivocal|maybe|might|may\s+have|"
    r"cannot\s+(?:exclude|rule\s+out)|not\s+excluded)\b",
    re.IGNORECASE,
)
SPLIT_RE = re.compile(r"[;\n]|,(?=\s)|\band\b", re.IGNORECASE)
MIN_LEXICAL_CHARS = 12
FAMILY_HISTORY_RE = re.compile(
    r"family\s+history|\bmaternal\b|\bpaternal\b|\bthe mother\b|\bmother'?s\b|"
    r"\bthe father\b|\bfather'?s\b|\bmiscarriage|\bpregnancy\s+loss",
    re.IGNORECASE,
)
GROUPING_NAME_RE = re.compile(r"^abnormality of\b", re.IGNORECASE)


@dataclass
class PhenotypeCandidate:
    hpo_id: str
    hpo_label: str
    status: str
    mapping_method: str
    confidence: str
    scope: str = SCOPE_PROBAND
    onset: str = ""
    frequency_context: str = ""
    source_phrase: str = ""
    obsolete: bool = False
    in_ontology: bool = True


@dataclass
class DocxStructure:
    paragraph_count: int = 0
    nonempty_paragraph_count: int = 0
    table_count: int = 0
    table_shapes: list[dict[str, int]] = field(default_factory=list)
    explicit_hpo_count: int = 0
    distinct_explicit_hpo: int = 0
    explicit_hpo_ids: list[str] = field(default_factory=list)
    explicit_unknown: list[str] = field(default_factory=list)
    explicit_obsolete: list[str] = field(default_factory=list)
    appears_structured: bool = False
    text_block_count: int = 0


def find_clinical_docx(data_dir: Path) -> Path:
    if not data_dir.is_dir():
        raise FileNotFoundError("dataset directory not found")
    matches = [
        path
        for path in data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() == ".docx"
    ]
    if not matches:
        raise FileNotFoundError("no DOCX found in the dataset directory")
    matches.sort(key=lambda item: item.stat().st_size, reverse=True)
    return matches[0]


def _iter_docx_blocks(path: Path) -> tuple[DocxStructure, list[str]]:
    from docx import Document

    document = Document(str(path))
    structure = DocxStructure()
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        structure.paragraph_count += 1
        text = (paragraph.text or "").strip()
        if text:
            structure.nonempty_paragraph_count += 1
            blocks.append(text)
    for table in document.tables:
        structure.table_count += 1
        rows = table.rows
        n_rows = len(rows)
        n_cols = len(rows[0].cells) if rows else 0
        structure.table_shapes.append({"rows": n_rows, "cols": n_cols})
        for row in rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            nonempty = [cell for cell in cells if cell]
            if nonempty:
                blocks.append(" | ".join(nonempty))
    structure.text_block_count = len(blocks)
    table_blocks = max(0, structure.text_block_count - structure.nonempty_paragraph_count)
    widest = max((shape["cols"] for shape in structure.table_shapes), default=0)
    structure.appears_structured = structure.table_count > 0 and (
        table_blocks >= structure.nonempty_paragraph_count or widest >= 2
    )
    return structure, blocks


def inspect_docx_structure(path: Path, ontology: HpoOntology | None = None) -> DocxStructure:
    structure, blocks = _iter_docx_blocks(path)
    found: list[str] = []
    for block in blocks:
        found.extend(HPO_ID_RE.findall(block))
    unique = list(dict.fromkeys(found))
    structure.explicit_hpo_count = len(found)
    structure.distinct_explicit_hpo = len(unique)
    structure.explicit_hpo_ids = unique
    if ontology is not None:
        unknown: list[str] = []
        obsolete: list[str] = []
        for hpo_id in unique:
            primary = ontology.primary_id(hpo_id)
            if primary is None:
                unknown.append(hpo_id)
            elif ontology.is_obsolete(primary):
                obsolete.append(hpo_id)
        structure.explicit_unknown = unknown
        structure.explicit_obsolete = obsolete
    return structure


def classify_status(text: str) -> str:
    has_neg = bool(NEGATION_RE.search(text))
    has_unc = bool(UNCERTAIN_RE.search(text))
    if has_neg and has_unc:
        return STATUS_UNCERTAIN
    if has_neg:
        return STATUS_ABSENT
    if has_unc:
        return STATUS_UNCERTAIN
    return STATUS_PRESENT


def _lookup_piece(ontology: HpoOntology, piece: str) -> list[tuple[str, str]]:
    hits = lookup_normalized(ontology, piece)
    if hits:
        return hits
    stripped = NEGATION_RE.sub(" ", piece)
    stripped = " ".join(stripped.split())
    if stripped and stripped.casefold() != piece.casefold():
        return lookup_normalized(ontology, stripped)
    return []


def classify_scope(text: str) -> str:
    if FAMILY_HISTORY_RE.search(text or ""):
        return SCOPE_FAMILY
    return SCOPE_PROBAND


def is_grouping_term(ontology: HpoOntology, hpo_id: str) -> bool:
    if hpo_id in GENERIC_HPO_IDS:
        return True
    return bool(GROUPING_NAME_RE.match(ontology.label(hpo_id)))


def _confidence_for_method(method: str, synonym_scope: str = "") -> str:
    if method == METHOD_EXPLICIT:
        return CONF_HIGH
    if method == METHOD_EXACT_LABEL:
        return CONF_HIGH
    if method == METHOD_SYNONYM:
        if synonym_scope in {"EXACT", "NARROW"}:
            return CONF_HIGH
        return CONF_MEDIUM
    return CONF_LOW


def _scope_to_method(scope: str) -> str:
    if scope == "EXACT_LABEL":
        return METHOD_EXACT_LABEL
    return METHOD_SYNONYM


def _phrase_pieces(block: str) -> list[str]:
    pieces = [block.strip()]
    for part in SPLIT_RE.split(block):
        text = part.strip(" |-")
        if text and text not in pieces:
            pieces.append(text)
    return pieces


def map_block(block: str, ontology: HpoOntology) -> list[PhenotypeCandidate]:
    block_scope = classify_scope(block)
    found: list[PhenotypeCandidate] = []
    seen: set[tuple[str, str, str]] = set()

    def _add(
        hpo_id: str,
        method: str,
        confidence: str,
        phrase: str,
        status: str,
        scope: str,
        *,
        allow_generic: bool = False,
    ) -> None:
        primary = ontology.primary_id(hpo_id) or hpo_id
        if primary in GENERIC_HPO_IDS and not allow_generic and method != METHOD_EXPLICIT:
            return
        if (
            status == STATUS_ABSENT
            and is_grouping_term(ontology, primary)
            and method != METHOD_EXPLICIT
        ):
            # Do not generalize a specific negation to a broad parent category.
            return
        key = (primary, status, scope)
        if key in seen:
            return
        term = ontology.terms.get(primary)
        in_ont = term is not None
        found.append(
            PhenotypeCandidate(
                hpo_id=primary,
                hpo_label=term.name if term else "",
                status=status,
                mapping_method=method,
                confidence=confidence,
                scope=scope,
                source_phrase=phrase,
                obsolete=bool(term and term.obsolete),
                in_ontology=in_ont,
            )
        )
        seen.add(key)

    pieces = _phrase_pieces(block)
    for piece in pieces:
        piece_status = classify_status(piece)
        piece_scope = classify_scope(piece)
        scope = piece_scope if piece_scope != SCOPE_PROBAND else block_scope
        for match in HPO_ID_RE.findall(piece):
            _add(
                match,
                METHOD_EXPLICIT,
                CONF_HIGH,
                block,
                piece_status,
                scope,
                allow_generic=True,
            )
        hits = _lookup_piece(ontology, piece)
        for hpo_id, syn_scope in hits:
            method = _scope_to_method(syn_scope)
            _add(
                hpo_id,
                method,
                _confidence_for_method(method, syn_scope),
                block,
                piece_status,
                scope,
            )

    # Whole-word lexical candidates only; never auto-promoted to the profile.
    norm_block = f" {normalize_label(block)} "
    if len(normalize_label(block)) >= MIN_LEXICAL_CHARS:
        block_status = classify_status(block)
        for hpo_id, term in ontology.terms.items():
            if term.obsolete or hpo_id in GENERIC_HPO_IDS:
                continue
            label_n = normalize_label(term.name)
            if len(label_n) < MIN_LEXICAL_CHARS:
                continue
            needle = f" {label_n} "
            if needle in norm_block and (hpo_id, block_status, block_scope) not in seen:
                _add(
                    hpo_id,
                    METHOD_LEXICAL,
                    CONF_LOW,
                    block,
                    block_status,
                    block_scope,
                )

    return _prefer_specific(found, ontology)


def _prefer_specific(
    candidates: list[PhenotypeCandidate],
    ontology: HpoOntology,
) -> list[PhenotypeCandidate]:
    ids = {item.hpo_id for item in candidates if item.mapping_method != METHOD_LEXICAL}
    drop: set[str] = set()
    for hid in ids:
        ancestors = ontology.ancestors(hid, include_self=False)
        for other in ids:
            if other != hid and other in ancestors:
                drop.add(other)
    kept: list[PhenotypeCandidate] = []
    for item in candidates:
        if item.mapping_method != METHOD_LEXICAL and item.hpo_id in drop:
            continue
        kept.append(item)
    return kept


def extract_phenotype(
    path: Path,
    ontology: HpoOntology,
) -> tuple[DocxStructure, list[PhenotypeCandidate]]:
    structure = inspect_docx_structure(path, ontology)
    _struct, blocks = _iter_docx_blocks(path)
    candidates: list[PhenotypeCandidate] = []
    seen: set[tuple[str, str, str, str]] = set()
    for block in blocks:
        for item in map_block(block, ontology):
            key = (item.hpo_id, item.status, item.mapping_method, item.scope)
            if key in seen:
                continue
            seen.add(key)
            candidates.append(item)
    return structure, drop_parent_when_child_present(candidates, ontology)


def drop_parent_when_child_present(
    candidates: list[PhenotypeCandidate],
    ontology: HpoOntology,
) -> list[PhenotypeCandidate]:
    present_ids = {
        item.hpo_id
        for item in candidates
        if item.status == STATUS_PRESENT and item.mapping_method != METHOD_LEXICAL
    }
    covered_ancestors: set[str] = set()
    for hid in present_ids:
        covered_ancestors |= set(ontology.ancestors(hid, include_self=False))
    kept: list[PhenotypeCandidate] = []
    for item in candidates:
        if item.mapping_method != METHOD_LEXICAL and item.hpo_id in covered_ancestors:
            continue
        kept.append(item)
    return kept


def profile_terms(
    candidates: list[PhenotypeCandidate],
    *,
    include_medium: bool = False,
) -> list[PhenotypeCandidate]:
    """Ranking profile: present terms at high (optionally medium) confidence."""
    allowed_conf = {CONF_HIGH} if not include_medium else {CONF_HIGH, CONF_MEDIUM}
    profile: list[PhenotypeCandidate] = []
    for item in candidates:
        if item.status != STATUS_PRESENT:
            continue
        if item.scope != SCOPE_PROBAND:
            continue
        if item.confidence not in allowed_conf:
            continue
        if item.mapping_method == METHOD_LEXICAL:
            continue
        if not item.in_ontology or item.obsolete:
            continue
        if item.hpo_id in GENERIC_HPO_IDS:
            continue
        profile.append(item)
    return profile


def safe_rows(candidates: list[PhenotypeCandidate]) -> list[dict[str, str]]:
    return [
        {
            "hpo_id": item.hpo_id,
            "hpo_label": item.hpo_label,
            "status": item.status,
            "confidence": item.confidence,
            "mapping_method": item.mapping_method,
            "scope": item.scope,
        }
        for item in candidates
        if item.mapping_method != METHOD_LEXICAL or item.confidence == CONF_LOW
    ]


def write_private_candidates(path: Path, candidates: list[PhenotypeCandidate]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = [
        "hpo_id",
        "hpo_label",
        "status",
        "mapping_method",
        "confidence",
        "scope",
        "onset",
        "frequency_context",
        "source_phrase",
        "obsolete",
        "in_ontology",
    ]
    lines = ["\t".join(columns)]
    for item in candidates:
        payload = asdict(item)
        values = []
        for col in columns:
            values.append(str(payload[col]).replace("\t", " ").replace("\n", " "))
        lines.append("\t".join(values))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_safe_phenotype(
    tsv_path: Path, json_path: Path, candidates: list[PhenotypeCandidate]
) -> None:
    rows = [
        row
        for row in safe_rows(candidates)
        if row["mapping_method"] != METHOD_LEXICAL
    ]
    tsv_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    header = ["hpo_id", "hpo_label", "status", "confidence", "mapping_method", "scope"]
    lines = ["\t".join(header)]
    for row in rows:
        lines.append("\t".join(row[col] for col in header))
    tsv_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    json_path.write_text(json.dumps(rows, indent=2) + "\n", encoding="utf-8")


def format_structure_report(structure: DocxStructure) -> str:
    return (
        "clinical_document_structure\n"
        f"paragraph_count={structure.paragraph_count}\n"
        f"nonempty_paragraph_count={structure.nonempty_paragraph_count}\n"
        f"table_count={structure.table_count}\n"
        f"table_shapes={structure.table_shapes}\n"
        f"text_block_count={structure.text_block_count}\n"
        f"appears_structured={structure.appears_structured}\n"
        f"explicit_hpo_count={structure.explicit_hpo_count}\n"
        f"distinct_explicit_hpo={structure.distinct_explicit_hpo}\n"
        f"explicit_hpo_ids={structure.explicit_hpo_ids}\n"
        f"explicit_unknown={structure.explicit_unknown}\n"
        f"explicit_obsolete={structure.explicit_obsolete}\n"
    )


def format_safe_phenotype_report(candidates: list[PhenotypeCandidate]) -> str:
    groups = {
        "high_confidence_present": [],
        "family_history": [],
        "high_confidence_absent": [],
        "uncertain": [],
        "medium_confidence": [],
        "manual_review_low_confidence": [],
    }
    for item in candidates:
        if item.mapping_method == METHOD_LEXICAL or item.confidence == CONF_LOW:
            groups["manual_review_low_confidence"].append(item)
            continue
        if item.scope == SCOPE_FAMILY:
            groups["family_history"].append(item)
            continue
        if item.status == STATUS_UNCERTAIN:
            groups["uncertain"].append(item)
        elif item.confidence == CONF_MEDIUM:
            groups["medium_confidence"].append(item)
        elif item.status == STATUS_ABSENT and item.confidence == CONF_HIGH:
            groups["high_confidence_absent"].append(item)
        elif item.status == STATUS_PRESENT and item.confidence == CONF_HIGH:
            groups["high_confidence_present"].append(item)
        else:
            groups["manual_review_low_confidence"].append(item)

    def _lines(label: str, rows: list[PhenotypeCandidate]) -> list[str]:
        out = [f"{label} n={len(rows)}"]
        for item in rows:
            out.append(
                f"  {item.hpo_id}\t{item.hpo_label}\t{item.status}\t"
                f"{item.confidence}\t{item.mapping_method}\t{item.scope}"
            )
        return out

    lines = ["phenotype_profile (HPO ids/labels only; no source text)"]
    for key in groups:
        lines.extend(_lines(key, groups[key]))
    return "\n".join(lines) + "\n"


def structure_as_dict(structure: DocxStructure) -> dict[str, Any]:
    return asdict(structure)
